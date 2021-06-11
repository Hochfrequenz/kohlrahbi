import re
from pathlib import Path
from typing import List

import docx
import pandas as pd
import xlsxwriter
from docx.table import Table
from docx.text.paragraph import Paragraph

from check_row_type import RowType, define_row_type
from get_sections import iter_block_items
from write_functions import write_new_row_in_dataframe

directory_path = Path.cwd() / "documents"
file_name = "UTILMD_AHB_WiM_3_1c_2021_04_01_2021_03_30.docx"

output_directory_path = Path.cwd() / "output"
json_output_directory_path = output_directory_path / "json"
csv_output_directory_path = output_directory_path / "csv"
xlsx_output_directory_path = output_directory_path / "xlsx"

json_output_directory_path.mkdir(parents=True, exist_ok=True)
csv_output_directory_path.mkdir(parents=True, exist_ok=True)
xlsx_output_directory_path.mkdir(parents=True, exist_ok=True)

path_to_file = directory_path / file_name

path_to_all_in_one_excel = xlsx_output_directory_path / f"{file_name[:-5]}.xlsx"

# Remove old all in one excel file if it exists
if path_to_all_in_one_excel.exists():
    path_to_all_in_one_excel.unlink(missing_ok=False)


def get_tabstop_positions(paragraph: Paragraph) -> List[int]:
    """Find all tabstop positions in a given paragraph

    Mainly the tabstop positions of cells from the middle column are determined

    Args:
        paragraph (Paragraph): Paragraph

    Returns:
        List[int]: All tabstop positions in the given paragraph
    """
    tabstop_positions: List = []
    for tabstop in paragraph.paragraph_format.tab_stops._pPr.tabs:
        tabstop_positions.append(tabstop.pos)
    return tabstop_positions


def read_table(
    table, dataframe, actual_df_row_index, last_two_row_types, left_indent_position: int, tabstop_positions: List
):

    if table._column_count == 4:
        index_for_middle_column = 2
    else:
        index_for_middle_column = 1

    for row in range(len(table.rows)):

        # initial empty list for the next row in the dataframe
        actual_dataframe_row = (len(dataframe.columns)) * [""]
        dataframe.loc[actual_df_row_index] = actual_dataframe_row

        row_cell_texts_as_list = [cell.text for cell in table.row_cells(row)]

        if table._column_count == 4:
            # remove redundant information for tables with 4 columns
            if (
                row_cell_texts_as_list[0] == row_cell_texts_as_list[1]
                and row_cell_texts_as_list[2] == row_cell_texts_as_list[3]
            ):
                # HEADER looks like
                # 0:'EDIFACT Struktur'
                # 1:'EDIFACT Struktur'
                # 2:'Beschreibung\tKündigung\tBestätigung\tAblehnung\tBedingung\n\tMSB \tKündigung\tKündigung\n\tMSB \tMSB \nKommunikation von\tMSBN an\tMSBA an\tMSBA an\n\tMSBA\tMSBN\tMSBN\nPrüfidentifikator\t11039\t11040\t11041'
                # 3:'Beschreibung\tKündigung\tBestätigung\tAblehnung\tBedingung\n\tMSB \tKündigung\tKündigung\n\tMSB \tMSB \nKommunikation von\tMSBN an\tMSBA an\tMSBA an\n\tMSBA\tMSBN\tMSBN\nPrüfidentifikator\t11039\t11040\t11041'
                # len():4
                del row_cell_texts_as_list[1]
                row_cell_texts_as_list[2] = ""
            elif row_cell_texts_as_list[1] == row_cell_texts_as_list[2]:
                # Dataelement row with header in the table
                # 0:'SG2\tNAD\t3035'
                # 1:'SG2\tNAD\t3035'
                # 2:'MR\tNachrichtenempfänger\tX\tX\tX'
                # 3:''
                # len():4
                del row_cell_texts_as_list[1]
            elif row_cell_texts_as_list[0] == row_cell_texts_as_list[1]:
                del row_cell_texts_as_list[1]

        actual_edifact_struktur_cell = table.row_cells(row)[0]

        # check for row type
        actual_row_type = define_row_type(
            edifact_struktur_cell=actual_edifact_struktur_cell,
            left_indent_position=left_indent_position,
        )
        print(actual_row_type.name)
        print(actual_edifact_struktur_cell.text)

        # write actual row into dataframe
        if actual_row_type is RowType.EMPTY and last_two_row_types[0] is RowType.HEADER:
            actual_df_row_index = actual_df_row_index - 1
            actual_df_row_index = write_new_row_in_dataframe(
                row_type=last_two_row_types[1],
                table=table,
                row=row,
                index_for_middle_column=index_for_middle_column,
                dataframe=dataframe,
                dataframe_row_index=actual_df_row_index,
                dataframe_row=actual_dataframe_row,
                row_cell_texts_as_list=row_cell_texts_as_list,
                left_indent_position=left_indent_position,
                tabstop_positions=tabstop_positions,
            )

        else:
            actual_df_row_index = write_new_row_in_dataframe(
                row_type=actual_row_type,
                table=table,
                row=row,
                index_for_middle_column=index_for_middle_column,
                dataframe=dataframe,
                dataframe_row_index=actual_df_row_index,
                dataframe_row=actual_dataframe_row,
                row_cell_texts_as_list=row_cell_texts_as_list,
                left_indent_position=left_indent_position,
                tabstop_positions=tabstop_positions,
            )

        # remember last row type for empty cells
        last_two_row_types[1] = last_two_row_types[0]
        last_two_row_types[0] = actual_row_type

    return last_two_row_types, actual_df_row_index


def beautify_bedingungen(bedingung):

    if isinstance(bedingung, str):
        bedingung = bedingung.replace("\n", " ")
        matches = re.findall(r"\[\d*\]", bedingung)
        for match in matches[1:]:
            index = bedingung.find(match)
            bedingung = bedingung[:index] + "\n" + bedingung[index:]
    return bedingung


def main():
    try:
        doc = docx.Document(path_to_file)  # Creating word reader object.

        pruefidentifikatoren: List = []

        # Iterate through the whole word document
        for item in iter_block_items(parent=doc):

            # Check if there is just a text paragraph,
            if isinstance(item, Paragraph) and not "Heading" in item.style.name:
                continue

            # Check if the paragraph is a chapter or section title
            elif isinstance(item, Paragraph) and "Heading" in item.style.name:
                chapter_title = item.text

                # Stop iterating at the section "Änderungshistorie"
                if chapter_title == "Änderungshistorie":
                    return 0

            # Check if a table comes with new Prüfidentifikatoren
            elif isinstance(item, Table) and item.cell(row_idx=0, col_idx=0).text == "EDIFACT Struktur":
                # check if list of pruefidentifikatoren is empty
                if pruefidentifikatoren:
                    # write for each pruefi an extra file
                    for pruefi in pruefidentifikatoren:
                        columns_to_export = base_columns + [pruefi]
                        columns_to_export.append("Bedingung")  # TODO: save string "Bedingung" in variable
                        df["Bedingung"] = df["Bedingung"].apply(beautify_bedingungen)
                        df_to_export = df[columns_to_export]
                        df_to_export.to_csv(csv_output_directory_path / f"{pruefi}.csv")

                        # for orient in ["split", "records", "index", "columns", "values", "table"]:

                        #     df_to_export.to_json(
                        #         json_output_directory_path / f"{pruefi}-{orient}.json", force_ascii=False, orient=orient
                        #     )
                        df_to_export.to_json(
                            json_output_directory_path / f"{pruefi}.json", force_ascii=False, orient="records"
                        )

                        with pd.ExcelWriter(
                            xlsx_output_directory_path / f"{pruefi}.xlsx", engine="xlsxwriter"
                        ) as writer:
                            df_to_export.to_excel(writer, sheet_name=f"{pruefi}")
                            workbook = writer.book
                            worksheet = writer.sheets[f"{pruefi}"]
                            wrap_format = workbook.add_format({"text_wrap": True})
                            column_letters = ["A", "B", "C", "D", "E", "F", "G", "H"]
                            column_widths = [3.5, 47, 9, 14, 39, 33, 18, 102]
                            for column_letter, column_width in zip(column_letters, column_widths):
                                excel_header = f"{column_letter}:{column_letter}"
                                worksheet.set_column(excel_header, column_width, wrap_format)

                        try:
                            with pd.ExcelWriter(path=path_to_all_in_one_excel, mode="a", engine="openpyxl") as writer:
                                df_to_export.to_excel(writer, sheet_name=f"{pruefi}")
                        except FileNotFoundError:
                            with pd.ExcelWriter(path=path_to_all_in_one_excel, mode="w", engine="openpyxl") as writer:
                                df_to_export.to_excel(writer, sheet_name=f"{pruefi}")

                print(chapter_title)
                print([cell.text for cell in item.row_cells(0)])
                print("\n NEW TABLE \n")

                header_cells = [cell.text for cell in item.row_cells(0)]
                look_up_term = "Prüfidentifikator"
                cutter_index = header_cells[-1].find(look_up_term) + 1
                # +1 cause of \t after Prüfidentifikator
                pruefidentifikatoren: List = header_cells[-1][cutter_index + len(look_up_term) :].split("\t")

                indicator_paragraph = item.cell(row_idx=4, col_idx=1).paragraphs[0]

                left_indent_position = indicator_paragraph.paragraph_format.left_indent
                tabstop_positions: List = get_tabstop_positions(indicator_paragraph)

                base_columns: List = [
                    "Segment Gruppe",
                    "Segment",
                    "Datenelement",
                    "Codes und Qualifier",
                    "Beschreibung",
                ]
                columns = base_columns + pruefidentifikatoren
                columns.append("Bedingung")

                df = pd.DataFrame(
                    columns=columns,
                    dtype="str",
                )
                # Initialize help variables
                last_two_row_types: List = [RowType.EMPTY, RowType.EMPTY]
                actual_df_row_index: int = 0

                last_two_row_types, actual_df_row_index = read_table(
                    table=item,
                    dataframe=df,
                    actual_df_row_index=actual_df_row_index,
                    last_two_row_types=last_two_row_types,
                    left_indent_position=left_indent_position,
                    tabstop_positions=tabstop_positions,
                )

            elif isinstance(item, Table) and "df" in locals():
                last_two_row_types, actual_df_row_index = read_table(
                    table=item,
                    dataframe=df,
                    actual_df_row_index=actual_df_row_index,
                    last_two_row_types=last_two_row_types,
                    left_indent_position=left_indent_position,
                    tabstop_positions=tabstop_positions,
                )

    except IOError:
        print("There was an error opening the file!")
        return


if __name__ == "__main__":
    main()
