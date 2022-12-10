import pandas as pd  # type:ignore[import]
from docx.table import _Cell  # type:ignore[import]


def parse_edifact_struktur_cell(
    table_cell: _Cell,
    dataframe: pd.DataFrame,
    row_index: int,
    edifact_struktur_cell_left_indent_position: int,
) -> None:
    """Parses a paragraph in the edifact struktur column and puts the information into the appropriate columns

    Args:
        table_cell (Cell): edifact struktur cell
        dataframe (pd.DataFrame): Contains all infos
        row_index (int): Current index of the DataFrame
        edifact_struktur_cell_left_indent_position (int): Position of the left indent from the indicator edifact
            struktur cell
    """

    joined_text = " ".join(p.text for p in table_cell.paragraphs)
    splitted_text_at_tabs = joined_text.split("\t")
    tab_count = joined_text.count("\t")

    # Check if the line starts on the far left
    if table_cell.paragraphs[0].paragraph_format.left_indent != edifact_struktur_cell_left_indent_position:

        if tab_count == 2:
            dataframe.at[row_index, "Segment Gruppe"] = splitted_text_at_tabs[0]
            dataframe.at[row_index, "Segment"] = splitted_text_at_tabs[1]
            dataframe.at[row_index, "Datenelement"] = splitted_text_at_tabs[2]
        elif tab_count == 1:
            dataframe.at[row_index, "Segment Gruppe"] = splitted_text_at_tabs[0]
            dataframe.at[row_index, "Segment"] = splitted_text_at_tabs[1]
        elif tab_count == 0 and joined_text.strip() != "":
            if table_cell.paragraphs[0].runs[0].bold:
                # Segmentgruppe: SG8
                dataframe.at[row_index, "Segment Gruppe"] = splitted_text_at_tabs[0]
            else:
                # Segmentname: Referenzen auf die ID der\nTranche
                _sg_text = dataframe.at[row_index, "Segment Gruppe"]
                if _sg_text == "":
                    # Referenzen auf die ID der
                    dataframe.at[row_index, "Segment Gruppe"] = splitted_text_at_tabs[0]
                else:
                    # Tranche
                    dataframe.at[row_index, "Segment Gruppe"] += " " + splitted_text_at_tabs[0]

    # Now the text should start in middle of the EDIFACT Struktur column
    else:

        if tab_count == 1:
            # Example: "UNH\t0062"
            dataframe.at[row_index, "Segment"] = splitted_text_at_tabs[0]
            dataframe.at[row_index, "Datenelement"] = splitted_text_at_tabs[1]

        elif tab_count == 0:
            # Example: "UNH"
            dataframe.at[row_index, "Segment"] = splitted_text_at_tabs[0]
